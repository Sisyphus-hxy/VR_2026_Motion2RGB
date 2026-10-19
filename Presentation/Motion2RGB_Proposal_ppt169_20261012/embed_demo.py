"""Enrich the official native SVG export with the original MP4 on slide 6."""
from pathlib import Path
from pptx import Presentation
from lxml import etree

root = Path(__file__).resolve().parent
prs = Presentation(root / "exports/Motion2RGB_base.pptx")
slide = prs.slides[5]
poster = next(shape for shape in slide.shapes if shape.name == "demo-video")
movie = slide.shapes.add_movie(
    str(root / "media/demo_video.mp4"),
    poster.left, poster.top, poster.width, poster.height,
    poster_frame_image=str(root / "images/video_08s.png"),
    mime_type="video/mp4",
)
movie.name = "Motion2RGB Demo - click to play (20.57 seconds)"
tree = poster._element.getparent()
index = tree.index(poster._element)
tree.remove(poster._element)
tree.remove(movie._element)
tree.insert(index, movie._element)

slide2_notes = prs.slides[1].notes_slide.notes_text_frame
slide2_notes.text += (
    "\n\n参考文献（不计入讲述）：\n"
    "Chattopadhyay et al. RobustNav. ICCV 2021.\n"
    "https://openaccess.thecvf.com/content/ICCV2021/html/"
    "Chattopadhyay_RobustNav_Towards_Benchmarking_Robustness_in_Embodied_Navigation_ICCV_2021_paper.html\n"
    "Tobin et al. Domain Randomization. IROS 2017.\n"
    "https://arxiv.org/abs/1703.06907"
)
destination = root / "exports/Motion2RGB_Proposal.pptx"
prs.save(destination)
print(destination)
